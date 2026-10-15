"""Tests für Booking-Import, Zusammenführung, Scoring und Dashboard (erfundene Beispieldaten)."""

import json

from hotelabgleich import booking, dashboard, demo, match, merge, places, report, scoring
from hotelabgleich.__main__ import build_ranking
from hotelabgleich.booking import BookingHotel
from hotelabgleich.members import Member
from hotelabgleich.merge import Hotel
from tests.test_hotelabgleich import hotel as ghotel


def h(name, g=None, gn=0, b=None, bn=0, **kw):
    return Hotel(id=name, name=name, google_bewertung=g, google_anzahl=gn, booking_bewertung=b, booking_anzahl=bn, **kw)


def test_bundesland_aus_plz():
    cases = {"1010": "Wien", "3500": "Niederösterreich", "4020": "Oberösterreich", "5280": "Oberösterreich",
             "5020": "Salzburg", "6020": "Tirol", "6900": "Vorarlberg", "9900": "Tirol", "7431": "Burgenland",
             "8010": "Steiermark", "9220": "Kärnten", "": ""}
    assert {p: merge.bundesland_aus_plz(p) for p in cases} == cases


def test_all_bundeslaender_have_tiles():
    assert len(places.BUNDESLAENDER) == 9


def test_google_amenities():
    p = {"allowsDogs": True, "goodForChildren": False, "parkingOptions": {"freeParkingLot": True}, "accessibilityOptions": {}}
    assert places.amenities(p) == ["Hunde erlaubt", "Parkplatz"]


def test_booking_from_api_and_subscores():
    item = {"id": 42, "name": {"de": "Hotel Beispiel"}, "accommodation_type": 204, "url": "https://www.booking.com/hotel/at/x.html",
            "location": {"address": {"de": "Dorf 1"}, "city": -1, "postal_code": "5600", "coordinates": {"latitude": 47.35, "longitude": 13.2}},
            "rating": {"stars": 4, "review_score": 8.9, "number_of_reviews": 321}, "facilities": [{"id": 2}, {"id": 999}]}
    b = booking.from_api(item, {"-1": "St. Johann im Pongau"}, {"2": "Parkplatz"}, {"204": "Hotel"})
    assert (b.name, b.ort, b.sterne, b.bewertung, b.anzahl_bewertungen, b.typ, b.ausstattung) == \
        ("Hotel Beispiel", "St. Johann im Pongau", 4.0, 8.9, 321, "Hotel", ["Parkplatz"])
    assert booking.parse_subscores({"data": [{"scores": {"cleanliness": {"score": 9.1}, "staff": 9.4}}]}) == {"Sauberkeit": 9.1, "Personal": 9.4}


def test_booking_csv(tmp_path):
    p = tmp_path / "b.csv"
    p.write_text("name;plz;bewertung;anzahl_bewertungen;ausstattung;sauberkeit\nHotel Test;6020;8,7;120;Spa|Pool;9,2\n", encoding="utf-8")
    b = booking.load_csv(p)[0]
    assert (b.bewertung, b.anzahl_bewertungen, b.ausstattung, b.teilnoten) == (8.7, 120, ["Spa", "Pool"], {"Sauberkeit": 9.2})


def test_merge_matches_nearby_same_name_and_keeps_rest():
    g = ghotel("Hotel Bergkristall", lat=47.8, lon=13.04)
    near = BookingHotel("1", "Bergkristall Hotel ****", plz="5020", lat=47.8003, lon=13.0401, bewertung=9.0, anzahl_bewertungen=50, sterne=4)
    other = BookingHotel("2", "Pension Waldesruh", plz="6020", lat=47.27, lon=11.40, bewertung=8.0, anzahl_bewertungen=10)
    out = merge.merge([g], [near, other])
    assert len(out) == 2
    assert (out[0].booking_bewertung, out[0].sterne, out[0].bundesland) == (9.0, 4, "Salzburg")
    assert (out[1].id, out[1].bundesland) == ("b:2", "Tirol")


def test_merge_does_not_join_different_names():
    g = ghotel("Hotel Sonnenhof", lat=47.8, lon=13.04)
    b = BookingHotel("1", "Hotel Waldesruh", lat=47.8001, lon=13.0401, bewertung=9.0, anzahl_bewertungen=50)
    assert len(merge.merge([g], [b])) == 2


def test_bayes_pulls_few_reviews_to_mean():
    few = h("Wenig", g=5.0, gn=3)
    many = h("Viel", g=4.8, gn=2000)
    avg = h("Mittel", g=4.0, gn=500)
    scoring.score_all([few, many, avg])
    assert many.score > few.score
    assert many.rang_gesamt == 1 and many.perzentil == 100.0
    assert few.vertrauen == "gering" and many.vertrauen == "hoch"


def test_score_combines_sources_and_flags_deviation():
    x = h("Beide", g=4.8, gn=1000, b=6.0, bn=1000)
    scoring.score_all([x, h("A", g=4.0, gn=10), h("B", b=8.0, bn=10)])
    assert x.score_google > 90 and x.score_booking < 60
    assert x.abweichung_quellen >= scoring.ABWEICHUNG_HINWEIS
    assert min(x.score_google, x.score_booking) < x.score < max(x.score_google, x.score_booking)


def test_no_reviews_no_score():
    x = h("Neu")
    scoring.score_all([x])
    assert x.score is None and x.rang_gesamt is None and x.klasse == "ohne Bewertungen"


def test_mann_whitney():
    assert scoring.mann_whitney_p([1, 2, 3], [1]) is None
    assert scoring.mann_whitney_p([90, 91, 92, 93, 94, 95], [60, 61, 62, 63, 64, 65]) < 0.01
    assert scoring.mann_whitney_p([80, 81, 82, 83], [80, 81, 82, 83]) > 0.9


def test_compare_adjusts_for_structure():
    # Mitglieder sind 5-Sterne-Häuser, die generell besser bewertet sind; innerhalb gleicher Sterne kein Unterschied.
    hs = []
    for i in range(6):
        hs.append(h(f"M5{i}", g=4.8, gn=300, sterne=5, bundesland="Tirol", mitglied_status="Mitglied"))
        hs.append(h(f"A5{i}", g=4.8, gn=300, sterne=5, bundesland="Tirol", mitglied_status="kein Mitglied"))
        hs.append(h(f"A3{i}", g=4.0, gn=300, sterne=3, bundesland="Tirol", mitglied_status="kein Mitglied"))
    hs.append(h("P", g=1.0, gn=300, mitglied_status="prüfen"))
    scoring.score_all(hs)
    cmp = scoring.compare(hs)
    assert cmp["gesamt"]["differenz"] > 5
    assert cmp["bereinigt"]["differenz"] == 0
    assert cmp["gesamt"]["andere"]["n"] == 12  # "prüfen" zählt nicht


def test_differing_website_is_not_member():
    hh = ghotel("Hotel Alpenrose", website="https://alpenrose-a.at")
    r = match.decide_by_rules(hh, match.candidates(hh, [Member("Hotel Alpenhof", plz="5020", website="alpenhof.at")]))
    assert (r.status, r.methode) == ("kein Mitglied", "Website abweichend")


def test_demo_pipeline_and_outputs(tmp_path):
    g, b, m = demo.generate(120)
    hotels = build_ranking(g, b, m, use_typesafe=False, log=lambda *_: None)
    assert any(x.ist_mitglied for x in hotels) and any(x.score for x in hotels)
    html = dashboard.render(hotels, "Demo", demo=True)
    data = json.loads(html.split('id="data">')[1].split("</script>")[0])
    assert data["demo"] and len(data["hotels"]) == len(hotels) and "bereinigt" in data["vergleich"]
    assert "<\\/" not in html.split('id="data">')[0]
    out = tmp_path / "r.xlsx"
    report.write_ranking(hotels, scoring.compare(hotels), out, "Demo")
    from openpyxl import load_workbook
    assert load_workbook(out).sheetnames == ["Vergleich", "Ranking", "Methodik"]
