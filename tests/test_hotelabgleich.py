"""Tests mit erfundenen Beispieldaten (keine echten Betriebe)."""

from types import SimpleNamespace

import pytest

from hotelabgleich import match, members, places, report
from hotelabgleich.members import Member
from hotelabgleich.places import GoogleHotel


def hotel(name, plz="5020", lat=47.80, lon=13.04, website="", **kw):
    base = dict(
        place_id=name, name=name, adresse=f"Teststraße 1, {plz} Testort", plz=plz, ort="Testort",
        bezirk="Salzburg", bundesland="Salzburg", lat=lat, lon=lon, website=website, telefon="",
        bewertung=4.5, anzahl_bewertungen=100, typ="hotel", status="OPERATIONAL", maps_url="",
    )
    return GoogleHotel(**{**base, **kw})


def test_normalize_name_strips_noise():
    assert match.normalize_name("Hotel Alpenblick GmbH ****S") == "alpenblick"
    assert match.normalize_name("Gasthof Zum Weißen Rößl") == "weissen rossl"


def test_domain():
    assert match.domain("https://www.beispielhotel.at/de/") == "beispielhotel.at"
    assert match.domain("beispielhotel.at") == "beispielhotel.at"
    assert match.domain("") == ""


def test_same_domain_is_member():
    h = hotel("Seehof Test", website="https://www.seehof-test.at/")
    ms = [Member("Hotel Seehof am See", plz="5330", website="seehof-test.at")]
    r = match.decide_by_rules(h, match.candidates(h, ms))
    assert r.status == "Mitglied" and r.methode == "gleiche Website-Domain"


def test_similar_name_same_plz_is_member():
    h = hotel("Hotel Bergkristall ****")
    r = match.decide_by_rules(h, match.candidates(h, [Member("Bergkristall Hotel GmbH", plz="5020")]))
    assert r.status == "Mitglied"


def test_far_away_namesake_is_not_candidate():
    h = hotel("Hotel Post", plz="5020")
    ms = [Member("Hotel Post", plz="5700", lat=47.32, lon=12.79)]
    assert match.decide_by_rules(h, match.candidates(h, ms)).status == "kein Mitglied"


def test_partial_name_goes_to_review():
    h = hotel("Alpenrose Apartments")
    r = match.decide_by_rules(h, match.candidates(h, [Member("Hotel Alpenrose", plz="5020")]))
    assert r.status == "prüfen"


def test_unrelated_neighbour_is_not_member():
    h = hotel("Hotel Sonnenhof")
    r = match.decide_by_rules(h, match.candidates(h, [Member("Hotel Waldesruh", plz="5020")]))
    assert r.status == "kein Mitglied"


def test_typesafe_resolves_review_cases(monkeypatch):
    import typesafe_sdk
    from typesafe_sdk._core.questions import normalize_questions

    seen = {}

    class FakeClient:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def system_one(self, state, questions):
            normalize_questions(questions)  # echte SDK-Validierung der Fragen
            seen["state"] = state
            return SimpleNamespace(nouls={"k0": SimpleNamespace(noul=0.93)})

    monkeypatch.setenv("TYPESAFE_API_KEY", "test")
    monkeypatch.setattr(typesafe_sdk, "TypeSafeClient", FakeClient)
    h = hotel("Alpenrose Apartments")
    results = [match.decide_by_rules(h, match.candidates(h, [Member("Hotel Alpenrose", plz="5020")]))]
    match.judge_with_typesafe(results, log=lambda *_: None)
    assert results[0].status == "Mitglied" and results[0].wahrscheinlichkeit == 0.93
    assert seen["state"]["kandidaten"][0]["name"] == "Hotel Alpenrose"


def test_typesafe_skipped_without_key(monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    h = hotel("Alpenrose Apartments")
    results = [match.decide_by_rules(h, match.candidates(h, [Member("Hotel Alpenrose", plz="5020")]))]
    match.judge_with_typesafe(results, log=lambda *_: None)
    assert results[0].status == "prüfen"


def test_parse_member_page_jsonld():
    html = """<script type="application/ld+json">{"@context":"https://schema.org","@type":"Hotel",
    "name":"Hotel Beispiel","url":"https://www.meinhotel.at/hotel-beispiel","sameAs":["https://www.hotel-beispiel.at"],
    "address":{"streetAddress":"Dorf 1","postalCode":"5600","addressLocality":"St. Johann"},
    "geo":{"latitude":"47.35","longitude":"13.2"}}</script>"""
    m = members.parse_member_page(html, "https://www.meinhotel.at/hotel-beispiel")
    assert (m.name, m.plz, m.website, m.lat) == ("Hotel Beispiel", "5600", "https://www.hotel-beispiel.at", 47.35)


def test_load_csv_semicolon(tmp_path):
    p = tmp_path / "m.csv"
    p.write_text("Name;PLZ;Ort;Website\nHotel Beispiel;5600;St. Johann;hotel-beispiel.at\n", encoding="utf-8")
    assert members.load_csv(p)[0].plz == "5600"


def test_google_from_api_and_split():
    p = {
        "id": "abc", "displayName": {"text": "Hotel Beispiel"}, "formattedAddress": "Dorf 1, 5600 St. Johann",
        "addressComponents": [
            {"longText": "5600", "types": ["postal_code"]},
            {"longText": "St. Johann im Pongau", "types": ["administrative_area_level_2", "political"]},
            {"longText": "Salzburg", "types": ["administrative_area_level_1", "political"]},
        ],
        "location": {"latitude": 47.35, "longitude": 13.2}, "userRatingCount": 12,
    }
    h = GoogleHotel.from_api(p)
    assert (h.plz, h.bezirk, h.bundesland) == ("5600", "St. Johann im Pongau", "Salzburg")
    assert len(places.split((0, 0, 2, 2))) == 4


def test_collect_subdivides_saturated_tiles():
    class FakeClient:
        def search_rect(self, query, rect, included_type):
            # Volle Kachel nur auf Ebene 0, danach je 1 Treffer pro Kachel
            if rect == places.BUNDESLAENDER["Salzburg"]:
                return [{"id": f"x{i}", "addressComponents": [{"longText": "Salzburg", "types": ["administrative_area_level_1"]}]} for i in range(60)]
            return [{"id": str(rect), "addressComponents": [{"longText": "Tirol", "types": ["administrative_area_level_1"]}]}]

    hs = places.collect(FakeClient(), "Salzburg", log=lambda *_: None)
    assert len(hs) == 60  # Tiroler Treffer aus den Teilkacheln werden verworfen


def test_report_writes(tmp_path):
    hs = [hotel("Hotel Bergkristall"), hotel("Hotel Sonnenhof")]
    results = match.match_all(hs, [Member("Bergkristall Hotel", plz="5020")], use_typesafe=False)
    out = tmp_path / "r.xlsx"
    report.write(results, 1, "Salzburg", out)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert wb.sheetnames == ["Übersicht", "Alle Betriebe", "Potenzial (Nicht-Mitglieder)", "Manuell prüfen", "Methodik"]
    assert wb["Übersicht"]["B5"].value == 1
