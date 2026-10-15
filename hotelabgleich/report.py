"""Excel-Bericht im ÖHV-Design (Arial, ÖHV-Blau als Tabellenkopf)."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .match import NOUL_MATCH, NOUL_NO_MATCH, Result

BLAU, BLAU_HELL, GELB = "007EAB", "DBEDF3", "EDA10A"
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BLAU)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=16, color=BLAU)
STATUS_FILL = {
    "Mitglied": PatternFill("solid", fgColor=BLAU_HELL),
    "prüfen": PatternFill("solid", fgColor="FCECC4"),
}

COLUMNS = [
    ("Status", 14), ("Betrieb (Google Maps)", 38), ("Bezirk", 16), ("PLZ", 7), ("Ort", 20),
    ("Adresse", 42), ("Website", 32), ("Telefon", 18), ("Google-Bewertung", 10), ("Anzahl Bewertungen", 11),
    ("Typ", 14), ("ÖHV-Mitglied (Treffer)", 36), ("Methode", 26), ("Jev-Wahrscheinlichkeit", 12),
    ("Mitglieds-Quelle", 36), ("Google-Maps-Link", 30),
]


def _row(r: Result) -> list:
    h, m = r.hotel, r.member
    return [
        r.status, h.name, h.bezirk, h.plz, h.ort, h.adresse, h.website, h.telefon, h.bewertung,
        h.anzahl_bewertungen, h.typ, m.name if m else "", r.methode,
        None if r.wahrscheinlichkeit is None else round(r.wahrscheinlichkeit, 3), m.quelle if m else "", h.maps_url,
    ]


def _table(ws, results: list[Result]) -> None:
    for c, (title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, title)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    for i, r in enumerate(results, 2):
        for c, v in enumerate(_row(r), 1):
            cell = ws.cell(i, c, v)
            cell.font = BODY_FONT
            if r.status in STATUS_FILL:
                cell.fill = STATUS_FILL[r.status]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def write(results: list[Result], members_total: int, bundesland: str, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Übersicht"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 14
    for col in "CDE":
        ws.column_dimensions[col].width = 14

    counts = Counter(r.status for r in results)
    total = len(results)
    ws["A1"] = f"ÖHV-Mitgliederabgleich {bundesland}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Stand {date.today():%d.%m.%Y} · Quellen: Google Places API, ÖHV-Mitgliederliste (meinhotel.at)"
    ws["A2"].font = Font(name="Arial", italic=True, size=9)

    rows = [
        ("Beherbergungsbetriebe auf Google Maps", total),
        ("davon ÖHV-Mitglied", counts["Mitglied"]),
        ("davon kein Mitglied (Potenzial)", counts["kein Mitglied"]),
        ("davon manuell zu prüfen", counts["prüfen"]),
        ("Mitgliedsquote (bestätigt)", counts["Mitglied"] / total if total else 0),
        ("Einträge in der Mitgliederliste (gesamt, alle Bundesländer)", members_total),
    ]
    for i, (label, val) in enumerate(rows, 4):
        ws.cell(i, 1, label).font = BODY_FONT
        cell = ws.cell(i, 2, val)
        cell.font = Font(name="Arial", bold=True, color=BLAU)
        if isinstance(val, float):
            cell.number_format = "0.0 %"

    start = 4 + len(rows) + 2
    for c, t in enumerate(["Bezirk", "Betriebe", "Mitglieder", "Potenzial", "Quote"], 1):
        cell = ws.cell(start, c, t)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    by_bezirk: dict[str, Counter] = defaultdict(Counter)
    for r in results:
        by_bezirk[r.hotel.bezirk or "(ohne Angabe)"][r.status] += 1
    for i, (bezirk, c) in enumerate(sorted(by_bezirk.items(), key=lambda kv: -sum(kv[1].values())), start + 1):
        n = sum(c.values())
        vals = [bezirk, n, c["Mitglied"], c["kein Mitglied"], c["Mitglied"] / n]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(i, col, v)
            cell.font = BODY_FONT
        ws.cell(i, 5).number_format = "0.0 %"

    order = {"Mitglied": 0, "prüfen": 1, "kein Mitglied": 2}
    _table(wb.create_sheet("Alle Betriebe"), sorted(results, key=lambda r: (order[r.status], r.hotel.bezirk, r.hotel.name)))
    potenzial = sorted((r for r in results if r.status == "kein Mitglied"), key=lambda r: -(r.hotel.anzahl_bewertungen or 0))
    _table(wb.create_sheet("Potenzial (Nicht-Mitglieder)"), potenzial)
    _table(wb.create_sheet("Manuell prüfen"), [r for r in results if r.status == "prüfen"])

    m = wb.create_sheet("Methodik")
    m.column_dimensions["A"].width = 120
    text = [
        "Methodik",
        "1. Google Places API (Text Search, Typ „lodging“, Suchbegriff „Hotel“): Das Bundesland wird in Kacheln zerlegt, "
        "volle Kacheln (60 Treffer) werden geviertelt. Nur Einträge mit Bundesland laut Google-Adresse, ohne dauerhaft geschlossene.",
        "2. Mitgliederliste: Hotelseiten von meinhotel.at (strukturierte Daten) bzw. CSV-Export der Mitgliederverwaltung.",
        "3. Zuordnung per Regel: gleiche Website-Domain, oder Namensähnlichkeit ≥ 90 % bei gleicher PLZ bzw. < 500 m Entfernung. "
        "Kein Mitglied, wenn alle namensähnlichen Mitglieder in der Nähe eine andere Website haben.",
        f"4. Unklare Fälle: TypeSafe-Modell Jev beurteilt je Kandidat, ob es derselbe Betrieb ist (Mitglied ab {NOUL_MATCH:.0%}, "
        f"ausgeschlossen unter {NOUL_NO_MATCH:.0%}, dazwischen manuelle Prüfung).",
        "",
        "Einschränkungen",
        "• Google Maps erfasst auch Pensionen, Apartments und Ferienwohnungen (Spalte „Typ“). Für Akquise-Listen ggf. nach Typ filtern.",
        "• Ein Mitglied kann mehreren Google-Einträgen zugeordnet sein (z. B. Hotel und Restaurant mit eigenem Eintrag).",
        "• Google-Daten sind nutzergeneriert; Bewertungszahl und Website sind Anhaltspunkte, keine geprüften Betriebsdaten.",
        "• Bewertungen und Kontaktdaten aus Google Maps unterliegen den Google-Maps-Nutzungsbedingungen (nur interne Auswertung, keine Weitergabe).",
    ]
    for i, t in enumerate(text, 1):
        cell = m.cell(i, 1, t)
        cell.font = Font(name="Arial", bold=t in ("Methodik", "Einschränkungen"), color=BLAU if t in ("Methodik", "Einschränkungen") else None)
        cell.alignment = Alignment(wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --- Ranking-Bericht (Gästescore) -----------------------------------------------------------------

RANK_COLUMNS = [
    ("Rang", 7), ("Rang Bundesland", 9), ("Betrieb", 38), ("Bundesland", 16), ("Bezirk", 18), ("PLZ", 7),
    ("Ort", 20), ("Typ", 14), ("Sterne", 7), ("ÖHV-Gästescore", 10), ("Klasse", 13), ("Vertrauen", 10),
    ("Perzentil", 9), ("Google", 8), ("Google Bew.", 10), ("Booking", 8), ("Booking Bew.", 10),
    ("Abweichung Google/Booking", 12), ("ÖHV-Status", 14), ("ÖHV-Mitglied (Treffer)", 32), ("Ausstattung", 40),
]


def _head(ws, row: int, titles: list[str]) -> None:
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row, c, t)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_ranking(hotels, cmp: dict, path: Path, quellen: str) -> None:
    from .scoring import KLASSEN, PRIOR_REVIEWS

    wb = Workbook()
    ws = wb.active
    ws.title = "Vergleich"
    ws["A1"] = "ÖHV-Gästescore: Mitglieder im Vergleich"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Stand {date.today():%d.%m.%Y} · Quellen: {quellen}"
    ws["A2"].font = Font(name="Arial", italic=True, size=9)
    for col, w in zip("ABCDEFGH", (28, 12, 12, 12, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = w

    titles = ["Gruppe", "Mitglieder n", "Ø Mitglieder", "Andere n", "Ø Andere", "Differenz", "p-Wert"]
    row = 4

    def block(title: str, groups: dict) -> None:
        nonlocal row
        ws.cell(row, 1, title).font = Font(name="Arial", bold=True, color=BLAU)
        row += 1
        _head(ws, row, titles)
        for name, g in groups.items():
            row += 1
            vals = [name, g["mitglieder"]["n"], g["mitglieder"]["mittel"], g["andere"]["n"], g["andere"]["mittel"],
                    g["differenz"], None if g["p_wert"] is None else round(g["p_wert"], 4)]
            for c, v in enumerate(vals, 1):
                ws.cell(row, c, v).font = BODY_FONT
        row += 2

    block("Gesamt", {"Alle Betriebe": cmp["gesamt"]})
    ws.cell(row, 1, "Strukturbereinigte Differenz (gleiches Bundesland, Sterne, Größe)").font = BODY_FONT
    ws.cell(row, 2, cmp["bereinigt"]["differenz"]).font = Font(name="Arial", bold=True, color=BLAU)
    ws.cell(row + 1, 1, "Mitgliedsanteil gesamt / Top 10 % / Flop 10 % (in %)").font = BODY_FONT
    for c, k in enumerate(("mitglieder_gesamt", "mitglieder_top10", "mitglieder_flop10"), 2):
        ws.cell(row + 1, c, cmp["anteile"][k]).font = BODY_FONT
    row += 3
    block("Nach Bundesland", cmp["bundesland"])
    block("Nach Sternen", cmp["sterne"])
    block("Nach Größe (Anzahl Bewertungen)", cmp["groesse"])

    rk = wb.create_sheet("Ranking")
    _head(rk, 1, [t for t, _ in RANK_COLUMNS])
    for c, (_, w) in enumerate(RANK_COLUMNS, 1):
        rk.column_dimensions[get_column_letter(c)].width = w
    scored = sorted((h for h in hotels if h.score is not None), key=lambda h: h.rang_gesamt)
    for i, h in enumerate(scored, 2):
        vals = [h.rang_gesamt, h.rang_bundesland, h.name, h.bundesland, h.bezirk, h.plz, h.ort, h.typ, h.sterne,
                h.score, h.klasse, h.vertrauen, h.perzentil, h.google_bewertung, h.google_anzahl or None,
                h.booking_bewertung, h.booking_anzahl or None, h.abweichung_quellen, h.mitglied_status,
                h.mitglied_name, ", ".join(h.ausstattung)]
        for c, v in enumerate(vals, 1):
            cell = rk.cell(i, c, v)
            cell.font = BODY_FONT
            if h.ist_mitglied:
                cell.fill = STATUS_FILL["Mitglied"]
    rk.freeze_panes = "D2"
    rk.auto_filter.ref = rk.dimensions

    m = wb.create_sheet("Methodik")
    m.column_dimensions["A"].width = 120
    text = [
        "Methodik ÖHV-Gästescore",
        "1. Normierung auf 0–100: Google (1–5 Sterne) → (x − 1) / 4 · 100; Booking.com (Note 1–10) → (x − 1) / 9 · 100.",
        f"2. Glättung je Quelle (Bayes-Mittel) mit {PRIOR_REVIEWS} fiktiven Durchschnittsbewertungen: wenige Bewertungen ziehen zum Durchschnitt.",
        "3. Gesamtscore: gewichteter Mittelwert der Quellen, Gewicht = 50 % · n / (n + Prior) je Quelle.",
        "4. Klassen: " + ", ".join(f"{label} ab {limit}" for limit, label in KLASSEN) + ". Vertrauen: hoch ab 200, mittel ab 50 Bewertungen.",
        "5. Vergleich: Mitglieder vs. Nicht-Mitglieder (ohne „prüfen“), Mann-Whitney-U-Test; strukturbereinigt innerhalb gleicher Zellen "
        "aus Bundesland × Sterne × Größenklasse, gewichtet nach Mitgliederzahl je Zelle.",
        "",
        "Einschränkungen",
        "• Zusammenhang ist kein Kausalbeleg: Ein höherer Score der Mitglieder zeigt nicht, dass die Mitgliedschaft ihn bewirkt.",
        "• Sterne laut Booking.com (teils Selbstangabe); offizielle Klassifizierung in Österreich: WKO/Hotelstars Union.",
        "• Google- und Booking-Inhalte unterliegen den Plattformbedingungen: nur interne Auswertung, regelmäßig neu abrufen.",
    ]
    for i, t in enumerate(text, 1):
        bold = t in ("Methodik ÖHV-Gästescore", "Einschränkungen")
        cell = m.cell(i, 1, t)
        cell.font = Font(name="Arial", bold=bold, color=BLAU if bold else None)
        cell.alignment = Alignment(wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
